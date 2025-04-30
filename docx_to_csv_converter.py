import pandas as pd
from docx import Document
import os

def convert_docx_to_csv(docx_path, output_csv_path):
    """
    Convert a DOCX file containing tables to a CSV file.
    """
    print(f"Converting {docx_path} to CSV...")
    doc = Document(docx_path)
    all_data = []
    all_headers = set()
    
    # First pass: collect all possible headers
    for table_idx, table in enumerate(doc.tables):
        print(f"Scanning table {table_idx + 1} for headers...")
        
        if len(table.rows) > 0:
            # Get headers from the first row
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            all_headers.update(headers)
    
    all_headers = sorted(list(all_headers))
    print(f"Found {len(all_headers)} unique headers: {all_headers}")
    
    # Second pass: extract data with consistent columns
    for table_idx, table in enumerate(doc.tables):
        print(f"Processing table {table_idx + 1}...")
        
        if len(table.rows) > 0:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            header_to_idx = {header: idx for idx, header in enumerate(headers)}
            
            # Process each row (skip the header row)
            for row in table.rows[1:]:
                # Initialize row data with empty values for all headers
                row_data = {header: "" for header in all_headers}
                
                # Fill in values from the row
                for idx, cell in enumerate(row.cells):
                    if idx < len(headers):
                        header = headers[idx]
                        row_data[header] = cell.text.strip()
                
                # Only add rows that have some data
                if any(row_data.values()):
                    all_data.append(row_data)
    
    if all_data:
        df = pd.DataFrame(all_data)
        df.to_csv(output_csv_path, index=False)
        print(f"Successfully converted to {output_csv_path}")
        print(f"Found {len(df)} rows of data")
    else:
        print("No data found in the document")

if __name__ == "__main__":
    docx_path = "elife-56879-supp3-v2.docx"
    output_csv_path = "patient_metadata.csv"
    
    convert_docx_to_csv(docx_path, output_csv_path) 